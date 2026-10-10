import math
from pathlib import Path

from openpyxl import load_workbook
from scipy.stats import gaussian_kde


def analizar_excel(ruta):
	libro = load_workbook(Path(ruta), data_only=True, read_only=True)
	datos = []
	for hoja in libro.worksheets:
		for fila in hoja.iter_rows(values_only=True):
			datos.extend(valor for valor in fila if isinstance(valor, (int, float)))

	if not datos:
		raise ValueError('El archivo no contiene datos numericos.')

	cantidad = len(datos)
	media = sum(datos) / cantidad
	varianza = sum((valor - media) ** 2 for valor in datos) / max(cantidad - 1, 1)
	desviacion = math.sqrt(varianza)
	minimo, maximo = min(datos), max(datos)

	numero_barras = min(12, max(5, math.ceil(math.sqrt(cantidad))))
	ancho = (maximo - minimo) / numero_barras if maximo != minimo else 1
	limites = [minimo + ancho * indice for indice in range(numero_barras + 1)]
	frecuencias = [0] * numero_barras
	for valor in datos:
		indice = min(int((valor - minimo) / ancho), numero_barras - 1)
		frecuencias[indice] += 1

	puntos_densidad = []
	if len(datos) > 1 and desviacion > 0:
		densidad = gaussian_kde(datos)
		paso = (maximo - minimo) / 79 if maximo != minimo else 1
		puntos = [minimo + paso * indice for indice in range(80)]
		puntos_densidad = [
			{'x': round(x, 6), 'y': round(float(y), 8)}
			for x, y in zip(puntos, densidad(puntos))
		]

	ordenados = sorted(datos)
	mediana = percentil(ordenados, 0.5)
	cuartil_1 = percentil(ordenados, 0.25)
	cuartil_3 = percentil(ordenados, 0.75)
	rango_intercuartil = cuartil_3 - cuartil_1
	limite_inferior = cuartil_1 - 1.5 * rango_intercuartil
	limite_superior = cuartil_3 + 1.5 * rango_intercuartil

	return {
		'datos': datos,
		'cantidad': cantidad,
		'media': media,
		'desviacion_estandar': desviacion,
		'minimo': minimo,
		'maximo': maximo,
		'histograma': {'limites': limites, 'frecuencias': frecuencias},
		'densidad': puntos_densidad,
		'boxplot': {
			'minimo': max(minimo, limite_inferior),
			'cuartil_1': cuartil_1,
			'mediana': mediana,
			'cuartil_3': cuartil_3,
			'maximo': min(maximo, limite_superior),
			'atipicos': [valor for valor in datos if valor < limite_inferior or valor > limite_superior],
		},
	}


def percentil(datos_ordenados, proporcion):
	posicion = (len(datos_ordenados) - 1) * proporcion
	inferior = math.floor(posicion)
	superior = math.ceil(posicion)
	if inferior == superior:
		return datos_ordenados[inferior]
	fraccion = posicion - inferior
	return datos_ordenados[inferior] + (datos_ordenados[superior] - datos_ordenados[inferior]) * fraccion